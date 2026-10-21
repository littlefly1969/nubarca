"""DOCX / PPTX / XLSX writers for the golden documents in office formats.

The files exist so that a NubArca-side extraction run can be compared against
the same ground truth. Office containers embed timestamps in their properties
and zip entries; both are pinned so regeneration is byte-identical.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Emu
from pptx.util import Inches as PInches
from pptx.util import Pt as PPt

from .model import C, Doc, F, H, L, N, P, T, as_cell
from .render import chart_png
from .truth import ordered_blocks

FIXED_TIME = dt.datetime(2025, 7, 1, 9, 0, 0)
AUTHOR = "NubArca document-bench (synthetic)"


def normalize_zip(path: Path) -> None:
    """Rewrite a zip with fixed entry timestamps, keeping entry order."""
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            target.writestr(info, data)
    path.write_bytes(buffer.getvalue())


def _grid(table: T) -> tuple[int, list[tuple[int, int, object]]]:
    """(columns, [(row, col, cell)]) with rowspan-covered slots skipped."""
    ncols = sum(as_cell(c).colspan for c in table.rows[0])
    covered: set[tuple[int, int]] = set()
    placed = []
    for r, row in enumerate(table.rows):
        c0 = 0
        for raw in row:
            while (r, c0) in covered:
                c0 += 1
            cell = as_cell(raw)
            placed.append((r, c0, cell))
            for dr in range(1, cell.rowspan):
                for dc in range(cell.colspan):
                    covered.add((r + dr, c0 + dc))
            c0 += cell.colspan
    return ncols, placed


# ---- DOCX -----------------------------------------------------------------

def write_docx(doc: Doc, path: Path) -> None:
    document = Document()
    props = document.core_properties
    props.author = AUTHOR
    props.title = doc.title
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.last_modified_by = AUTHOR
    props.revision = 1
    for index, page in enumerate(doc.pages):
        if index:
            document.add_page_break()
        for block in ordered_blocks(doc, page):
            if isinstance(block, H):
                document.add_heading(block.text, level=block.level)
            elif isinstance(block, P):
                document.add_paragraph(block.text)
            elif isinstance(block, L):
                for item in block.items:
                    document.add_paragraph(item, style="List Bullet")
            elif isinstance(block, T):
                if block.caption:
                    document.add_paragraph(block.caption, style="Caption")
                ncols, placed = _grid(block)
                table = document.add_table(rows=len(block.rows), cols=ncols)
                table.style = "Table Grid"
                for r, c0, cell in placed:
                    target = table.cell(r, c0)
                    if cell.colspan > 1 or cell.rowspan > 1:
                        target = target.merge(table.cell(r + cell.rowspan - 1, c0 + cell.colspan - 1))
                    target.text = cell.text
                    if r < block.header_rows:
                        for run in target.paragraphs[0].runs:
                            run.bold = True
            elif isinstance(block, F):
                document.add_paragraph(block.plain)
            elif isinstance(block, C):
                document.add_picture(io.BytesIO(chart_png(block)), width=Inches(6))
                if block.caption:
                    document.add_paragraph(block.caption, style="Caption")
            elif isinstance(block, N):
                note = document.add_paragraph(block.text)
                for run in note.runs:
                    run.font.size = Pt(8)
    document.save(str(path))
    normalize_zip(path)


# ---- PPTX -----------------------------------------------------------------

def write_pptx(doc: Doc, path: Path) -> None:
    deck = Presentation()
    deck.slide_width = PInches(13.333)
    deck.slide_height = PInches(7.5)
    deck.core_properties.author = AUTHOR
    deck.core_properties.title = doc.title
    deck.core_properties.created = FIXED_TIME
    deck.core_properties.modified = FIXED_TIME
    deck.core_properties.last_modified_by = AUTHOR
    deck.core_properties.revision = 1
    blank = deck.slide_layouts[6]
    width = deck.slide_width - PInches(1.0)
    for page in doc.pages:
        slide = deck.slides.add_slide(blank)
        top = PInches(0.5)
        for block in ordered_blocks(doc, page):
            if isinstance(block, H):
                box = slide.shapes.add_textbox(PInches(0.5), top, width, PInches(0.9))
                box.text_frame.text = block.text
                box.text_frame.paragraphs[0].runs[0].font.size = PPt(32)
                box.text_frame.paragraphs[0].runs[0].font.bold = True
                top += PInches(1.0)
            elif isinstance(block, (P, L)):
                lines = [block.text] if isinstance(block, P) else list(block.items)
                box = slide.shapes.add_textbox(PInches(0.5), top, width, PInches(0.5 * len(lines)))
                frame = box.text_frame
                frame.word_wrap = True
                for i, line in enumerate(lines):
                    para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    para.text = ("• " + line) if isinstance(block, L) else line
                    para.runs[0].font.size = PPt(20)
                top += PInches(0.55 * len(lines) + 0.2)
            elif isinstance(block, T):
                ncols, placed = _grid(block)
                shape = slide.shapes.add_table(len(block.rows), ncols, PInches(0.5), top, width,
                                               PInches(0.45 * len(block.rows)))
                table = shape.table
                for r, c0, cell in placed:
                    target = table.cell(r, c0)
                    if cell.colspan > 1 or cell.rowspan > 1:
                        target.merge(table.cell(r + cell.rowspan - 1, c0 + cell.colspan - 1))
                    target.text = cell.text
                top += PInches(0.45 * len(block.rows) + 0.3)
            elif isinstance(block, C):
                picture = slide.shapes.add_picture(io.BytesIO(chart_png(block)), PInches(2.5), top,
                                                   height=PInches(4.6))
                top += Emu(picture.height) + PInches(0.1)
                if block.caption:
                    box = slide.shapes.add_textbox(PInches(0.5), top, width, PInches(0.5))
                    box.text_frame.text = block.caption
                    top += PInches(0.6)
        if page.notes:
            slide.notes_slide.notes_text_frame.text = page.notes
    deck.save(str(path))
    normalize_zip(path)


# ---- XLSX -----------------------------------------------------------------

def write_xlsx(doc: Doc, path: Path) -> None:
    workbook = Workbook()
    workbook.properties.creator = AUTHOR
    workbook.properties.title = doc.title
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    workbook.properties.lastModifiedBy = AUTHOR
    workbook.remove(workbook.active)
    for index, page in enumerate(doc.pages):
        sheet = workbook.create_sheet(doc.sheet_names[index])
        row_offset = 0
        for block in page.blocks:
            if not isinstance(block, T):
                continue
            _, placed = _grid(block)
            for r, c0, cell in placed:
                target = sheet.cell(row=row_offset + r + 1, column=c0 + 1, value=cell.text)
                if r < block.header_rows:
                    target.font = Font(bold=True)
                if cell.colspan > 1 or cell.rowspan > 1:
                    sheet.merge_cells(start_row=row_offset + r + 1, start_column=c0 + 1,
                                      end_row=row_offset + r + cell.rowspan,
                                      end_column=c0 + cell.colspan)
            row_offset += len(block.rows) + 1
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 26
    workbook.save(str(path))
    normalize_zip(path)


WRITERS = {"docx": write_docx, "pptx": write_pptx, "xlsx": write_xlsx}
